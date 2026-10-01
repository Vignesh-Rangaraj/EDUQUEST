from __future__ import annotations

"""Revise manuscript AI claims and replace its Fig. 5 image from verified source."""

import argparse
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


OLD_TO_NEW = {
    "curriculum-integrated AI features built on the Google Gemini API":
    "a lesson-aware, multilingual AI Tutor built on the Google Gemini API",
    "curriculum-integrated AI features": "a lesson-aware Gemini AI Tutor",
    "curriculum- and progress-integrated AI features with explicitly specified offline degradation behavior":
    "a lesson-aware, multilingual AI Tutor with structured modes and an explicit online dependency",
}

FIGURE_CAPTION = (
    "Fig. 5.  Implemented AI Tutor workflow: authenticated student request, optional lesson context, "
    "Google Gemini response, chat persistence, and browser-native voice support."
)
FIGURE_PROSE = (
    "Fig. 5 summarizes the implemented EduQuest AI Tutor. An authenticated student submits a message "
    "in CHAT, EXPLAIN, SIMPLIFY, HINT, or PRACTICE mode. The backend validates the message and selected "
    "language; when a lesson identifier is supplied, it retrieves the Activity title and description and "
    "associated LessonContent, then includes that material as primary context in the Gemini prompt. "
    "The Gemini service uses backend-managed configuration and handles provider timeouts and transient "
    "failures. Successful replies and user messages are stored in the chat session. The Tutor supports "
    "English, Tamil, Hindi, Malayalam, Telugu, and Kannada responses. PRACTICE mode can request two or "
    "three practice questions as part of the conversation; it does not create or publish a teacher-authored "
    "quiz. Browser speech recognition supports English, Tamil, and Hindi input, while browser speech "
    "synthesis reads replies using the selected language. These voice capabilities use browser APIs, not "
    "Gemini audio generation."
)
OFFLINE_PROSE = (
    "AI Tutor responses require an online backend connection and access to Google Gemini. When the client "
    "is offline, the Tutor interface reports that internet connectivity is required and disables voice "
    "features; new Gemini prompts are not queued for later delivery. This network dependency is limited to "
    "the AI Tutor and does not change the platform's offline lesson, quiz, or progress flows."
)


def replace_paragraph(paragraph, value: str) -> None:
    paragraph.clear()
    paragraph.add_run(value)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("source", type=Path)
    parser.add_argument("output", type=Path)
    parser.add_argument("figure", type=Path)
    args = parser.parse_args()
    doc = Document(args.source)

    changed = 0
    for p in doc.paragraphs:
        original = p.text
        if original.strip() == "VII. AI MODULE":
            replace_paragraph(p, "VII. AI TUTOR")
            changed += 1
        value = original
        for old, new in OLD_TO_NEW.items():
            value = value.replace(old, new)
        if value != original:
            replace_paragraph(p, value)
            changed += 1
        if original.startswith("Fig. 5."):
            replace_paragraph(p, FIGURE_CAPTION)
            changed += 1
        elif original.startswith("Fig. 5 (placeholder"):
            replace_paragraph(p, FIGURE_PROSE)
            changed += 1
        elif original.startswith("Because all four features depend on network access"):
            replace_paragraph(p, OFFLINE_PROSE)
            changed += 1

    table_updates = 0
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip() == "Curriculum-integrated":
                    replace_paragraph(cell.paragraphs[0], "Lesson-aware Tutor (optional context)")
                    table_updates += 1
                if cell.text.strip() == "Quiz gen., tutor, summarization, analytics assistant":
                    replace_paragraph(cell.paragraphs[0], "AI Tutor; five modes; optional lesson context; six languages")
                    table_updates += 1

    paragraphs = doc.paragraphs
    caption_index = next((i for i, p in enumerate(paragraphs) if p.text.startswith("Fig. 5.")), None)
    if caption_index is None:
        raise RuntimeError("Could not locate Fig. 5 caption")
    figure_paragraph = next(
        (p for p in reversed(paragraphs[:caption_index]) if p._p.xpath(".//a:blip")),
        None,
    )
    if figure_paragraph is None:
        raise RuntimeError("Could not locate the existing Fig. 5 image paragraph")
    figure_paragraph.clear()
    figure_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    shape = figure_paragraph.add_run().add_picture(str(args.figure), width=Inches(3.35))
    shape._inline.docPr.set("descr", "Implemented EduQuest AI Tutor workflow with optional lesson context, Gemini, chat persistence and browser voice support")

    if changed < 5 or table_updates < 2:
        raise RuntimeError(f"Expected manuscript changes were not all found: paragraphs={changed}, table cells={table_updates}")
    args.output.parent.mkdir(parents=True, exist_ok=True)
    doc.save(args.output)
    print(f"Saved {args.output}; updated {changed} paragraphs, {table_updates} table cells, and Fig. 5.")


if __name__ == "__main__":
    main()
